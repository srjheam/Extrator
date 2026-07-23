"""Baixa, valida e publica as bases Qualis oficiais da Plataforma Sucupira."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys
import tempfile
import time
import unicodedata
import warnings
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from email.message import Message
from pathlib import Path
from typing import Callable, Iterable, Sequence
from urllib.parse import urljoin
from xml.etree import ElementTree

import requests
from bs4 import BeautifulSoup
from bs4.element import Tag
from openpyxl import load_workbook


AREA_SUPORTADA = "COMPUTAÇÃO"
QUADRIENIO_SUPORTADO = "2021-2024"
POLITICA_IMPORTACAO_VERSAO = "1"
ESTRATOS_VALIDOS = frozenset(
    {"A1", "A2", "A3", "A4", "B1", "B2", "B3", "B4", "C"}
)
MIN_EVENTOS = 700
MIN_PERIODICOS = 2_500

COLUNAS_EVENTOS = ("Sigla", "Nome do evento", "Estrato")
COLUNAS_PERIODICOS = ("ISSN", "Título", "Estrato")

URL_EVENTOS = (
    "https://sucupira-legado.capes.gov.br/sucupira/public/consultas/"
    "coleta/qualisEventos/listaQualisEventos.xhtml"
)
URL_PERIODICOS = (
    "https://sucupira-legado.capes.gov.br/sucupira/public/consultas/"
    "coleta/veiculoPublicacaoQualis/listaConsultaGeralPeriodicos.xhtml"
)
OPCAO_EVENTOS = "Classificação de trabalho em anais 2025"
OPCAO_PERIODICOS = "CLASSIFICAÇÕES DE PERIÓDICOS QUADRIÊNIO 2021-2024"

CONTENT_TYPES_XLSX = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
        "application/zip",
    }
)


class ImportacaoError(RuntimeError):
    """Erro que impede a publicação das bases."""


@dataclass(frozen=True)
class Fonte:
    tipo: str
    url: str
    opcao: str
    adiciona_area: bool = False


@dataclass(frozen=True)
class Download:
    conteudo: bytes
    nome_original: str
    opcao_selecionada: str


@dataclass(frozen=True)
class DadosNormalizados:
    linhas: tuple[dict[str, str], ...]
    contagem_antes: int
    contagem_depois: int


FONTES = {
    "eventos": Fonte("eventos", URL_EVENTOS, OPCAO_EVENTOS),
    "periodicos": Fonte(
        "periodicos",
        URL_PERIODICOS,
        OPCAO_PERIODICOS,
        adiciona_area=True,
    ),
}


def _texto(valor: object) -> str:
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFC", str(valor))
    return re.sub(r"\s+", " ", texto).strip()


def _chave_texto(valor: object) -> str:
    return unicodedata.normalize("NFKC", _texto(valor)).casefold()


def validar_alvo(area: str, quadrienio: str) -> None:
    if _chave_texto(area) != _chave_texto(AREA_SUPORTADA):
        raise ImportacaoError(
            f"Área não suportada: {area!r}. Use {AREA_SUPORTADA!r}."
        )
    if _texto(quadrienio) != QUADRIENIO_SUPORTADO:
        raise ImportacaoError(
            "Quadriênio não suportado: "
            f"{quadrienio!r}. Use {QUADRIENIO_SUPORTADO!r}."
        )


def _sha256(conteudo: bytes) -> str:
    return hashlib.sha256(conteudo).hexdigest()


def _agora_utc() -> str:
    return (
        datetime.now(timezone.utc)
        .replace(microsecond=0)
        .isoformat()
        .replace("+00:00", "Z")
    )


def _formulario(html: str, pagina_url: str) -> tuple[BeautifulSoup, Tag, str, str]:
    soup = BeautifulSoup(html, "html.parser")
    form = soup.find("form", id="form")
    if not isinstance(form, Tag):
        raise ImportacaoError("A página Sucupira não contém o formulário principal.")
    viewstate = form.find("input", attrs={"name": "javax.faces.ViewState"})
    if not isinstance(viewstate, Tag) or not viewstate.get("value"):
        raise ImportacaoError("O formulário Sucupira não contém javax.faces.ViewState.")
    action = urljoin(pagina_url, str(form.get("action") or pagina_url))
    return soup, form, action, str(viewstate["value"])


def _resolver_opcao(form: Tag, nome_select: str, descricao: str) -> tuple[str, str]:
    select = form.find("select", attrs={"name": nome_select})
    if not isinstance(select, Tag):
        raise ImportacaoError(f"Campo JSF ausente: {nome_select}.")
    opcoes = [
        option
        for option in select.find_all("option")
        if _chave_texto(option.get_text(" ", strip=True)) == _chave_texto(descricao)
    ]
    if len(opcoes) != 1 or not opcoes[0].get("value"):
        raise ImportacaoError(
            f"Opção {descricao!r} não foi encontrada de forma única em {nome_select}."
        )
    return str(opcoes[0]["value"]), _texto(opcoes[0].get_text(" ", strip=True))


def _resposta_ajax(xml: str) -> dict[str, str]:
    try:
        raiz = ElementTree.fromstring(xml)
    except ElementTree.ParseError as exc:
        raise ImportacaoError("A resposta AJAX da Sucupira não é XML válido.") from exc
    erro = raiz.find(".//error")
    if erro is not None:
        mensagem = _texto(" ".join(erro.itertext()))
        raise ImportacaoError(f"A Sucupira retornou erro AJAX: {mensagem}")
    return {
        str(update.attrib.get("id")): update.text or ""
        for update in raiz.findall(".//update")
        if update.attrib.get("id")
    }


def _confirmar_opcao(
    form: Tag,
    nome_select: str,
    valor: str,
    descricao: str,
) -> None:
    select = form.find("select", attrs={"name": nome_select})
    if not isinstance(select, Tag):
        raise ImportacaoError(f"Campo JSF ausente após a consulta: {nome_select}.")
    selecionadas = select.find_all("option", selected=True)
    if len(selecionadas) != 1:
        raise ImportacaoError(f"A Sucupira não confirmou a seleção de {descricao!r}.")
    selecionada = selecionadas[0]
    if (
        str(selecionada.get("value")) != valor
        or _chave_texto(selecionada.get_text(" ", strip=True))
        != _chave_texto(descricao)
    ):
        raise ImportacaoError(f"A Sucupira selecionou uma opção diferente de {descricao!r}.")


def _controle_area(form: Tag) -> tuple[Tag, Tag]:
    area = form.find("select", attrs={"name": "form:area"})
    if not isinstance(area, Tag):
        raise ImportacaoError("Campo JSF da área não foi encontrado.")
    grupo = area.find_parent(class_="input-group")
    check = grupo.find("input", attrs={"type": "checkbox"}) if grupo else None
    if not isinstance(check, Tag) or not check.get("name"):
        raise ImportacaoError("Checkbox JSF da área não foi encontrado.")
    return area, check


def _controle_consultar(form: Tag) -> Tag:
    controles = [
        controle
        for controle in form.find_all("input", attrs={"type": "submit"})
        if _chave_texto(controle.get("value")) == _chave_texto("Consultar")
    ]
    if len(controles) != 1 or not controles[0].get("name"):
        raise ImportacaoError("Ação JSF Consultar não foi encontrada.")
    return controles[0]


def _acao_ajax_area(form: Tag) -> tuple[str, str, str, str]:
    area, _ = _controle_area(form)
    for link in form.find_all("a", id=True):
        onclick = str(link.get("onclick") or "")
        match = re.search(
            r"mojarra\.ab\(this,event,'([^']*)','([^']*)','([^']*)'\)",
            onclick,
        )
        if match and str(area.get("name")) in match.group(2).split():
            comportamento, executar, renderizar = match.groups()
            return str(link["id"]), comportamento, executar, renderizar
    raise ImportacaoError("Ação AJAX para adicionar a área não foi encontrada.")


def _comando_download(link: Tag) -> tuple[str, str]:
    onclick = str(link.get("onclick") or "")
    match = re.search(r"\{\s*'([^']+)'\s*:\s*'([^']+)'\s*\}", onclick)
    if not match:
        raise ImportacaoError("A ação JSF do arquivo oficial não foi encontrada.")
    return match.group(1), match.group(2)


def _link_arquivo_oficial(soup: BeautifulSoup, area: str) -> Tag:
    candidatos: list[Tag] = []
    for imagem in soup.find_all("img"):
        descricao = imagem.get("alt") or imagem.get("title") or ""
        if _chave_texto(descricao) != _chave_texto("Arquivo de classificações"):
            continue
        linha = imagem.find_parent("tr")
        if isinstance(linha, Tag) and _chave_texto(area) in _chave_texto(
            linha.get_text(" ", strip=True)
        ):
            link = imagem.find_parent("a")
            if isinstance(link, Tag):
                candidatos.append(link)
    if len(candidatos) != 1:
        raise ImportacaoError(
            "O link oficial da área selecionada não foi encontrado de forma única."
        )
    return candidatos[0]


def _campos_formulario(form: Tag) -> dict[str, str]:
    campos: dict[str, str] = {}
    form_id = str(form.get("id") or "")
    if form_id:
        campos[form_id] = form_id
    for controle in form.find_all(["input", "select", "textarea"]):
        nome = controle.get("name")
        if not nome:
            continue
        if controle.name == "input":
            tipo = str(controle.get("type") or "text").lower()
            if tipo in {"submit", "reset", "button", "image", "file"}:
                continue
            if tipo in {"checkbox", "radio"} and not controle.has_attr("checked"):
                continue
            campos[str(nome)] = str(controle.get("value") or "")
        elif controle.name == "select":
            opcao = controle.find("option", selected=True) or controle.find("option")
            if isinstance(opcao, Tag):
                campos[str(nome)] = str(opcao.get("value") or "")
        else:
            campos[str(nome)] = controle.get_text()
    return campos


def _nome_content_disposition(cabecalho: str | None, padrao: str) -> str:
    if not cabecalho:
        return padrao
    mensagem = Message()
    mensagem["content-disposition"] = cabecalho
    nome = mensagem.get_filename()
    return _texto(nome) if nome else padrao


def validar_xlsx(conteudo: bytes) -> None:
    if not conteudo.startswith(b"PK") or not zipfile.is_zipfile(io.BytesIO(conteudo)):
        raise ImportacaoError("O arquivo recebido não é um XLSX válido.")
    with zipfile.ZipFile(io.BytesIO(conteudo)) as arquivo:
        nomes = set(arquivo.namelist())
    obrigatorios = {"[Content_Types].xml", "xl/workbook.xml"}
    if not obrigatorios.issubset(nomes):
        raise ImportacaoError("O ZIP recebido não contém uma pasta de trabalho XLSX.")


def validar_resposta_xlsx(resposta: requests.Response) -> None:
    content_type = str(resposta.headers.get("Content-Type") or "")
    content_type = content_type.split(";", 1)[0].strip().lower()
    if content_type not in CONTENT_TYPES_XLSX:
        raise ImportacaoError(
            f"Content-Type incompatível com XLSX: {content_type or '(ausente)'}."
        )
    validar_xlsx(resposta.content)


class SucupiraClient:
    """Cliente mínimo para o fluxo JSF público da Sucupira."""

    def __init__(
        self,
        session: requests.Session | None = None,
        timeout: tuple[float, float] = (10, 90),
        tentativas: int = 3,
        sleep: Callable[[float], None] = time.sleep,
    ) -> None:
        self.session = session or requests.Session()
        self.timeout = timeout
        self.tentativas = tentativas
        self.sleep = sleep
        self.session.headers.update(
            {"User-Agent": "QualisLens-Sucupira/1 (+importacao-auditavel)"}
        )

    def _request(self, method: str, url: str, **kwargs: object) -> requests.Response:
        ultimo_status: int | None = None
        for tentativa in range(1, self.tentativas + 1):
            try:
                resposta = self.session.request(
                    method,
                    url,
                    timeout=self.timeout,
                    **kwargs,
                )
            except requests.Timeout as exc:
                if tentativa == self.tentativas:
                    raise ImportacaoError(
                        f"Timeout ao acessar a Sucupira após {tentativa} tentativas."
                    ) from exc
                self.sleep(0.5 * (2 ** (tentativa - 1)))
                continue
            except requests.RequestException as exc:
                raise ImportacaoError(f"Falha HTTP sem retry: {exc}") from exc

            ultimo_status = resposta.status_code
            retry = resposta.status_code == 429 or resposta.status_code >= 500
            if retry and tentativa < self.tentativas:
                self.sleep(0.5 * (2 ** (tentativa - 1)))
                continue
            if retry:
                raise ImportacaoError(
                    f"Sucupira retornou HTTP {resposta.status_code} "
                    f"após {tentativa} tentativas."
                )
            try:
                resposta.raise_for_status()
            except requests.RequestException as exc:
                raise ImportacaoError(
                    f"Sucupira retornou HTTP {resposta.status_code}; sem retry."
                ) from exc
            return resposta
        raise ImportacaoError(f"Falha HTTP inesperada: {ultimo_status}.")

    def baixar(self, fonte: Fonte, area_descricao: str) -> Download:
        inicial = self._request("GET", fonte.url)
        _, form, action, viewstate = _formulario(inicial.text, inicial.url)
        evento_valor, evento_texto = _resolver_opcao(
            form,
            "form:evento",
            fonte.opcao,
        )
        area_valor, area_texto = _resolver_opcao(
            form,
            "form:area",
            area_descricao,
        )

        form_id = str(form.get("id") or "form")
        evento_ajax = {
            "javax.faces.partial.ajax": "true",
            "javax.faces.source": "form:evento",
            "javax.faces.partial.execute": form_id,
            "javax.faces.partial.render": form_id,
            "javax.faces.behavior.event": "change",
            "javax.faces.partial.event": "change",
            form_id: form_id,
            "form:evento": evento_valor,
            "form:area": "0",
            "javax.faces.ViewState": viewstate,
        }
        resposta_ajax = self._request(
            "POST",
            action,
            data=evento_ajax,
            headers={"Faces-Request": "partial/ajax"},
        )
        updates = _resposta_ajax(resposta_ajax.text)
        if form_id not in updates or "javax.faces.ViewState" not in updates:
            raise ImportacaoError("A resposta AJAX não atualizou o formulário e o ViewState.")
        form_atualizado = BeautifulSoup(updates[form_id], "html.parser").find(
            "form",
            id=form_id,
        )
        if not isinstance(form_atualizado, Tag):
            raise ImportacaoError("O formulário atualizado pela Sucupira é inválido.")
        _confirmar_opcao(
            form_atualizado,
            "form:evento",
            evento_valor,
            evento_texto,
        )
        viewstate = updates["javax.faces.ViewState"]

        if fonte.adiciona_area:
            source, comportamento, executar, renderizar = _acao_ajax_area(
                form_atualizado
            )
            adicionar_area = {
                "javax.faces.partial.ajax": "true",
                "javax.faces.source": source,
                "javax.faces.partial.execute": f"{executar} {source}".strip(),
                "javax.faces.partial.render": renderizar,
                "javax.faces.behavior.event": comportamento,
                "javax.faces.partial.event": "click",
                form_id: form_id,
                source: source,
                "form:area": area_valor,
                "javax.faces.ViewState": viewstate,
            }
            resposta_area = self._request(
                "POST",
                action,
                data=adicionar_area,
                headers={"Faces-Request": "partial/ajax"},
            )
            updates_area = _resposta_ajax(resposta_area.text)
            if (
                renderizar not in updates_area
                or _chave_texto(area_texto)
                not in _chave_texto(
                    BeautifulSoup(
                        updates_area[renderizar],
                        "html.parser",
                    ).get_text(" ", strip=True)
                )
                or "javax.faces.ViewState" not in updates_area
            ):
                raise ImportacaoError("A Sucupira não confirmou a área adicionada.")
            viewstate = updates_area["javax.faces.ViewState"]

        area_select, area_check = _controle_area(form_atualizado)
        consultar = _controle_consultar(form_atualizado)
        consulta = {
            form_id: form_id,
            "form:evento": evento_valor,
            str(area_select["name"]): area_valor,
            str(area_check["name"]): str(area_check.get("value") or "on"),
            str(consultar["name"]): str(consultar.get("value") or "Consultar"),
            "javax.faces.ViewState": viewstate,
        }
        pagina_resultado = self._request("POST", action, data=consulta)
        soup_resultado, form_resultado, action_resultado, _ = _formulario(
            pagina_resultado.text,
            pagina_resultado.url,
        )
        _confirmar_opcao(
            form_resultado,
            "form:evento",
            evento_valor,
            evento_texto,
        )
        _confirmar_opcao(
            form_resultado,
            str(area_select["name"]),
            area_valor,
            area_texto,
        )
        link = _link_arquivo_oficial(soup_resultado, area_texto)
        comando, valor_comando = _comando_download(link)

        campos = _campos_formulario(form_resultado)
        campos.update(
            {
                "form:evento": evento_valor,
                str(area_select["name"]): area_valor,
                str(area_check["name"]): str(area_check.get("value") or "on"),
                comando: valor_comando,
            }
        )
        resposta_arquivo = self._request("POST", action_resultado, data=campos)
        validar_resposta_xlsx(resposta_arquivo)
        nome = _nome_content_disposition(
            resposta_arquivo.headers.get("Content-Disposition"),
            f"sucupira_{fonte.tipo}.xlsx",
        )
        return Download(resposta_arquivo.content, nome, evento_texto)


def _linhas_xlsx(
    conteudo: bytes,
    colunas: Sequence[str],
    tipo: str,
) -> list[dict[str, str]]:
    validar_xlsx(conteudo)
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style",
            module="openpyxl.styles.stylesheet",
        )
        try:
            workbook = load_workbook(
                io.BytesIO(conteudo),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except Exception as exc:
            raise ImportacaoError(f"Não foi possível ler o XLSX de {tipo}.") from exc
    try:
        if len(workbook.worksheets) != 1:
            raise ImportacaoError(
                f"XLSX de {tipo} deve conter exatamente uma planilha."
            )
        worksheet = workbook.worksheets[0]
        # A exportação oficial de periódicos declara dimension ref="A1".
        # As linhas existem. O recálculo torna todas elas visíveis no modo read-only.
        worksheet.reset_dimensions()
        iterator = worksheet.iter_rows(values_only=True)
        try:
            cabecalho_bruto = next(iterator)
        except StopIteration as exc:
            raise ImportacaoError(f"XLSX de {tipo} está vazio.") from exc
        cabecalho = list(cabecalho_bruto)
        while cabecalho and not _texto(cabecalho[-1]):
            cabecalho.pop()
        cabecalho_exato = tuple(
            "" if celula is None else str(celula) for celula in cabecalho
        )
        if cabecalho_exato != tuple(colunas):
            raise ImportacaoError(
                f"Colunas de {tipo} inválidas: {cabecalho_exato!r}; "
                f"esperado {tuple(colunas)!r}."
            )

        linhas: list[dict[str, str]] = []
        for numero, linha in enumerate(iterator, start=2):
            valores = list(linha)
            if not any(_texto(valor) for valor in valores):
                continue
            extras = valores[len(colunas) :]
            if any(_texto(valor) for valor in extras):
                raise ImportacaoError(
                    f"Linha {numero} de {tipo} contém coluna inesperada."
                )
            valores = valores[: len(colunas)]
            valores.extend([None] * (len(colunas) - len(valores)))
            linhas.append(
                {
                    coluna: _texto(valor)
                    for coluna, valor in zip(colunas, valores, strict=True)
                }
            )
        return linhas
    finally:
        workbook.close()


def _validar_campos_e_estratos(
    linhas: Iterable[dict[str, str]],
    colunas: Sequence[str],
    tipo: str,
) -> None:
    for numero, linha in enumerate(linhas, start=2):
        vazias = [coluna for coluna in colunas if not linha[coluna]]
        if vazias:
            raise ImportacaoError(
                f"Linha {numero} de {tipo} contém campos vazios: {vazias}."
            )
        if linha["Estrato"] not in ESTRATOS_VALIDOS:
            raise ImportacaoError(
                f"Estrato inválido na linha {numero} de {tipo}: "
                f"{linha['Estrato']!r}."
            )


def _validar_duplicatas_e_conflitos(
    linhas: Sequence[dict[str, str]],
    colunas: Sequence[str],
    identidade: Callable[[dict[str, str]], str],
    tipo: str,
) -> None:
    vistos: set[tuple[str, ...]] = set()
    estratos: dict[str, str] = {}
    for numero, linha in enumerate(linhas, start=2):
        registro = tuple(linha[coluna] for coluna in colunas)
        if registro in vistos:
            raise ImportacaoError(f"Duplicata exata na linha {numero} de {tipo}.")
        vistos.add(registro)
        chave = identidade(linha)
        anterior = estratos.get(chave)
        if anterior is not None and anterior != linha["Estrato"]:
            raise ImportacaoError(
                f"Identidade {chave!r} tem estratos conflitantes em {tipo}: "
                f"{anterior} e {linha['Estrato']}."
            )
        estratos[chave] = linha["Estrato"]


def issn_valido(issn: str) -> bool:
    if not re.fullmatch(r"\d{4}-\d{3}[\dX]", issn):
        return False
    digitos = issn.replace("-", "")
    soma = sum(int(digito) * peso for digito, peso in zip(digitos[:7], range(8, 1, -1)))
    verificador = (11 - soma % 11) % 11
    esperado = "X" if verificador == 10 else str(verificador)
    return digitos[-1] == esperado


def normalizar_eventos(
    conteudo: bytes,
    minimo: int = MIN_EVENTOS,
) -> DadosNormalizados:
    linhas = _linhas_xlsx(conteudo, COLUNAS_EVENTOS, "eventos")
    antes = len(linhas)
    normalizadas = [
        {
            "Sigla": _texto(linha["Sigla"]),
            "Nome do evento": _texto(linha["Nome do evento"]),
            "Estrato": _texto(linha["Estrato"]).upper(),
        }
        for linha in linhas
    ]
    _validar_campos_e_estratos(normalizadas, COLUNAS_EVENTOS, "eventos")
    _validar_duplicatas_e_conflitos(
        normalizadas,
        COLUNAS_EVENTOS,
        lambda linha: _chave_texto(linha["Sigla"]),
        "eventos",
    )
    if len(normalizadas) < minimo:
        raise ImportacaoError(
            f"Eventos abaixo do mínimo: {len(normalizadas)}; esperado ao menos {minimo}."
        )
    normalizadas.sort(
        key=lambda linha: (
            _chave_texto(linha["Sigla"]),
            _chave_texto(linha["Nome do evento"]),
            linha["Estrato"],
        )
    )
    return DadosNormalizados(tuple(normalizadas), antes, len(normalizadas))


def normalizar_periodicos(
    conteudo: bytes,
    minimo: int = MIN_PERIODICOS,
) -> DadosNormalizados:
    linhas = _linhas_xlsx(conteudo, COLUNAS_PERIODICOS, "periódicos")
    antes = len(linhas)
    normalizadas = [
        {
            "ISSN": _texto(linha["ISSN"]).upper(),
            "Título": _texto(linha["Título"]),
            "Estrato": _texto(linha["Estrato"]).upper(),
        }
        for linha in linhas
    ]
    _validar_campos_e_estratos(normalizadas, COLUNAS_PERIODICOS, "periódicos")
    for numero, linha in enumerate(normalizadas, start=2):
        if not issn_valido(linha["ISSN"]):
            raise ImportacaoError(
                f"ISSN inválido na linha {numero}: {linha['ISSN']!r}."
            )
    _validar_duplicatas_e_conflitos(
        normalizadas,
        COLUNAS_PERIODICOS,
        lambda linha: linha["ISSN"],
        "periódicos",
    )
    if len(normalizadas) < minimo:
        raise ImportacaoError(
            "Periódicos abaixo do mínimo: "
            f"{len(normalizadas)}; esperado ao menos {minimo}."
        )
    normalizadas.sort(
        key=lambda linha: (
            linha["ISSN"],
            _chave_texto(linha["Título"]),
            linha["Estrato"],
        )
    )
    return DadosNormalizados(tuple(normalizadas), antes, len(normalizadas))


def _csv_bytes(
    colunas: Sequence[str],
    linhas: Iterable[dict[str, str]],
) -> bytes:
    saida = io.StringIO(newline="")
    escritor = csv.DictWriter(
        saida,
        fieldnames=list(colunas),
        extrasaction="ignore",
        lineterminator="\n",
    )
    escritor.writeheader()
    escritor.writerows(linhas)
    return saida.getvalue().encode("utf-8")


def gerar_csvs(
    eventos: DadosNormalizados,
    periodicos: DadosNormalizados,
    quadrienio: str,
) -> tuple[bytes, bytes, bytes]:
    eventos_csv = _csv_bytes(
        (*COLUNAS_EVENTOS, "quadrienio"),
        ({**linha, "quadrienio": quadrienio} for linha in eventos.linhas),
    )
    periodicos_csv = _csv_bytes(
        (*COLUNAS_PERIODICOS, "quadrienio"),
        ({**linha, "quadrienio": quadrienio} for linha in periodicos.linhas),
    )
    projecao = _csv_bytes(
        ("ISSN", "Estrato"),
        (
            {"ISSN": linha["ISSN"], "Estrato": linha["Estrato"]}
            for linha in periodicos.linhas
        ),
    )
    return eventos_csv, periodicos_csv, projecao


def gerar_metadata(
    *,
    area: str,
    quadrienio: str,
    gerado_em_utc: str,
    modo: str,
    eventos_download: Download,
    periodicos_download: Download,
    eventos: DadosNormalizados,
    periodicos: DadosNormalizados,
    eventos_csv: bytes,
    periodicos_csv: bytes,
    projecao_csv: bytes,
) -> dict[str, object]:
    return {
        "area": area,
        "quadrienio": quadrienio,
        "gerado_em_utc": gerado_em_utc,
        "modo": modo,
        "politica_importacao_versao": POLITICA_IMPORTACAO_VERSAO,
        "fontes": {
            "eventos": {
                "url": URL_EVENTOS,
                "opcao_selecionada": eventos_download.opcao_selecionada,
                "nome_original": eventos_download.nome_original,
                "arquivo_original": (
                    "QualisLens/base/originais/"
                    "sucupira_eventos_computacao_2021_2024.xlsx"
                ),
                "sha256_xlsx": _sha256(eventos_download.conteudo),
                "arquivo_csv": "QualisLens/base/qualis_2021_2024.csv",
                "sha256_csv": _sha256(eventos_csv),
                "quantidade_registros": eventos.contagem_depois,
                "contagem_antes_normalizacao": eventos.contagem_antes,
                "contagem_depois_normalizacao": eventos.contagem_depois,
            },
            "periodicos": {
                "url": URL_PERIODICOS,
                "opcao_selecionada": periodicos_download.opcao_selecionada,
                "nome_original": periodicos_download.nome_original,
                "arquivo_original": (
                    "QualisLens/base/originais/"
                    "sucupira_periodicos_computacao_2021_2024.xlsx"
                ),
                "sha256_xlsx": _sha256(periodicos_download.conteudo),
                "arquivo_csv": (
                    "Classificador/base/qualis_periodicos_2021_2024.csv"
                ),
                "sha256_csv": _sha256(periodicos_csv),
                "quantidade_registros": periodicos.contagem_depois,
                "contagem_antes_normalizacao": periodicos.contagem_antes,
                "contagem_depois_normalizacao": periodicos.contagem_depois,
            },
        },
        "artefatos": {
            "qualis_unificado": {
                "arquivo": "Classificador/qualis-unificado.csv",
                "sha256_csv": _sha256(projecao_csv),
                "quantidade_registros": periodicos.contagem_depois,
            }
        },
    }


def _publicar_atomicamente(
    arquivos: Sequence[tuple[Path, Path]],
    staging: Path,
) -> None:
    for origem, _ in arquivos:
        if not origem.is_file():
            raise ImportacaoError(f"Artefato temporário ausente: {origem}.")
    for _, destino in arquivos:
        destino.parent.mkdir(parents=True, exist_ok=True)

    backups = staging / "backups"
    backups.mkdir()
    processados: list[tuple[Path, Path | None]] = []
    try:
        for indice, (origem, destino) in enumerate(arquivos):
            backup = backups / str(indice) if destino.exists() else None
            if backup is not None:
                os.replace(destino, backup)
            try:
                os.replace(origem, destino)
            except Exception:
                if backup is not None and backup.exists():
                    os.replace(backup, destino)
                raise
            processados.append((destino, backup))
    except Exception as exc:
        for destino, backup in reversed(processados):
            if backup is None:
                destino.unlink(missing_ok=True)
            else:
                destino.unlink(missing_ok=True)
                if backup.exists():
                    os.replace(backup, destino)
        raise ImportacaoError("Falha durante a publicação atômica.") from exc


def executar_importacao(
    *,
    area: str = AREA_SUPORTADA,
    quadrienio: str = QUADRIENIO_SUPORTADO,
    offline: bool = False,
    eventos_xlsx: Path | None = None,
    periodicos_xlsx: Path | None = None,
    repo_root: Path | None = None,
    client: SucupiraClient | None = None,
    gerado_em_utc: str | None = None,
    min_eventos: int = MIN_EVENTOS,
    min_periodicos: int = MIN_PERIODICOS,
) -> dict[str, object]:
    validar_alvo(area, quadrienio)
    area = AREA_SUPORTADA
    quadrienio = QUADRIENIO_SUPORTADO
    raiz = (repo_root or Path(__file__).resolve().parents[2]).resolve()

    if offline:
        if eventos_xlsx is None or periodicos_xlsx is None:
            raise ImportacaoError(
                "O modo offline exige --eventos-xlsx e --periodicos-xlsx."
            )
        try:
            eventos_conteudo = eventos_xlsx.read_bytes()
            periodicos_conteudo = periodicos_xlsx.read_bytes()
        except OSError as exc:
            raise ImportacaoError(f"Não foi possível ler os XLSX offline: {exc}") from exc
        eventos_download = Download(
            eventos_conteudo,
            eventos_xlsx.name,
            OPCAO_EVENTOS,
        )
        periodicos_download = Download(
            periodicos_conteudo,
            periodicos_xlsx.name,
            OPCAO_PERIODICOS,
        )
        modo = "offline"
    else:
        if eventos_xlsx is not None or periodicos_xlsx is not None:
            raise ImportacaoError(
                "Use --eventos-xlsx e --periodicos-xlsx somente com --offline."
            )
        http = client or SucupiraClient()
        eventos_download = http.baixar(FONTES["eventos"], area)
        periodicos_download = http.baixar(FONTES["periodicos"], area)
        modo = "online"

    eventos = normalizar_eventos(eventos_download.conteudo, min_eventos)
    periodicos = normalizar_periodicos(
        periodicos_download.conteudo,
        min_periodicos,
    )
    eventos_csv, periodicos_csv, projecao_csv = gerar_csvs(
        eventos,
        periodicos,
        quadrienio,
    )
    metadata = gerar_metadata(
        area=area,
        quadrienio=quadrienio,
        gerado_em_utc=gerado_em_utc or _agora_utc(),
        modo=modo,
        eventos_download=eventos_download,
        periodicos_download=periodicos_download,
        eventos=eventos,
        periodicos=periodicos,
        eventos_csv=eventos_csv,
        periodicos_csv=periodicos_csv,
        projecao_csv=projecao_csv,
    )
    metadata_bytes = (
        json.dumps(
            metadata,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )
        + "\n"
    ).encode("utf-8")

    raiz.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=".qualis-sucupira-",
        dir=raiz,
    ) as temporario:
        staging = Path(temporario)
        conteudos = (
            (
                "eventos.xlsx",
                eventos_download.conteudo,
                raiz
                / "QualisLens/base/originais/"
                "sucupira_eventos_computacao_2021_2024.xlsx",
            ),
            (
                "periodicos.xlsx",
                periodicos_download.conteudo,
                raiz
                / "QualisLens/base/originais/"
                "sucupira_periodicos_computacao_2021_2024.xlsx",
            ),
            (
                "eventos.csv",
                eventos_csv,
                raiz / "QualisLens/base/qualis_2021_2024.csv",
            ),
            (
                "periodicos.csv",
                periodicos_csv,
                raiz / "Classificador/base/qualis_periodicos_2021_2024.csv",
            ),
            (
                "qualis-unificado.csv",
                projecao_csv,
                raiz / "Classificador/qualis-unificado.csv",
            ),
            (
                "metadata.json",
                metadata_bytes,
                raiz / "QualisLens/base/metadata.json",
            ),
        )
        publicacao: list[tuple[Path, Path]] = []
        for nome, conteudo, destino in conteudos:
            origem = staging / nome
            origem.write_bytes(conteudo)
            publicacao.append((origem, destino))
        _publicar_atomicamente(publicacao, staging)
    return metadata


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Baixa, valida e publica as bases Qualis oficiais de Computação."
        )
    )
    parser.add_argument("--area", default=AREA_SUPORTADA)
    parser.add_argument("--quadrienio", default=QUADRIENIO_SUPORTADO)
    parser.add_argument("--offline", action="store_true")
    parser.add_argument("--eventos-xlsx", type=Path)
    parser.add_argument("--periodicos-xlsx", type=Path)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        metadata = executar_importacao(
            area=args.area,
            quadrienio=args.quadrienio,
            offline=args.offline,
            eventos_xlsx=args.eventos_xlsx,
            periodicos_xlsx=args.periodicos_xlsx,
        )
    except ImportacaoError as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        return 1
    eventos = metadata["fontes"]["eventos"]  # type: ignore[index]
    periodicos = metadata["fontes"]["periodicos"]  # type: ignore[index]
    print(
        "Publicação concluída: "
        f"{eventos['quantidade_registros']} eventos e "  # type: ignore[index]
        f"{periodicos['quantidade_registros']} periódicos."  # type: ignore[index]
    )
    print(
        "Normalização: "
        f"eventos {eventos['contagem_antes_normalizacao']} -> "  # type: ignore[index]
        f"{eventos['contagem_depois_normalizacao']}; "  # type: ignore[index]
        f"periódicos {periodicos['contagem_antes_normalizacao']} -> "  # type: ignore[index]
        f"{periodicos['contagem_depois_normalizacao']}."  # type: ignore[index]
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
